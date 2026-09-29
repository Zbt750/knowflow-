from pathlib import Path

import pytest
from pydantic import ValidationError

from backend.api.routes.chat import CreateSessionRequest, RenameSessionRequest, AskRequest
from backend.config import Settings
from backend.ingestion import document_parsers as parsers


@pytest.mark.parametrize("model,field", [(CreateSessionRequest, "title"), (RenameSessionRequest, "title"), (AskRequest, "question")])
def test_blank_chat_input_rejected(model, field):
    with pytest.raises(ValidationError):
        model(**{field: "   "})


def test_docx_block_order_preserved(tmp_path):
    from docx import Document
    document = Document()
    document.add_heading("第一章", 1)
    document.add_table(rows=1, cols=1).cell(0, 0).text = "第一章表格"
    document.add_heading("第二章", 1)
    path = tmp_path / "ordered.docx"
    document.save(path)
    result = parsers.parse_docx(path)
    assert result.parser_version == "docx-v2"
    assert result.normalized_text.index("第一章表格") < result.normalized_text.index("第二章")


def test_docx_expanded_text_rejected(tmp_path, monkeypatch):
    from docx import Document
    document = Document()
    document.add_paragraph("x" * 1000)
    path = tmp_path / "large.docx"
    document.save(path)
    monkeypatch.setattr(parsers, "MAX_DOCUMENT_CHARS", 100)
    with pytest.raises(parsers.DocumentParseError, match="document_content_too_large"):
        parsers.parse_docx(path)


def test_test_storage_defaults_not_development():
    settings = Settings(_env_file=None, app_env="test", database_url="postgresql://u:p@localhost/dev", test_database_url="postgresql://u:p@localhost/isolated_test")
    assert settings.upload_dir.resolve() != Path("storage/uploads").resolve()
    assert settings.chroma_dir.resolve() != Path("storage/chroma").resolve()


def test_untrusted_origin_cannot_mutate(app, monkeypatch):
    from fastapi.testclient import TestClient
    monkeypatch.setattr("backend.app._try_build_retrieval_stack", lambda *args: False)
    with TestClient(app) as client:
        response = client.post("/api/chat/sessions", json={"title": "blocked"}, headers={"Origin": "https://untrusted.example"})
        assert response.status_code == 403
        assert response.json()["error"]["code"] == "invalid_request"
        assert client.post("/api/chat/sessions", json={"title": "allowed"}, headers={"Origin": "http://localhost:5173"}).status_code == 201


def test_production_host_is_not_implicitly_an_allowed_origin(app, monkeypatch):
    """A reverse-proxy Host value must not extend the mutation Origin allowlist."""
    from fastapi.testclient import TestClient

    monkeypatch.setattr("backend.app._try_build_retrieval_stack", lambda *args: False)
    with TestClient(app) as client:
        app.state.settings = app.state.settings.model_copy(update={"app_env": "prod"})
        response = client.post(
            "/api/chat/sessions",
            json={"title": "host-origin-bypass"},
            headers={"Host": "attacker.example", "Origin": "http://attacker.example"},
        )
        assert response.status_code == 403
        assert response.json()["error"]["code"] == "invalid_request"


def test_lost_instance_lock_returns_safe_503(app, monkeypatch):
    from fastapi.testclient import TestClient

    monkeypatch.setattr("backend.app._try_build_retrieval_stack", lambda *args: False)
    with TestClient(app) as client:
        app.state.instance_lock_lost = True
        response = client.get("/api/health")
        assert response.status_code == 503
        assert response.json()["error"]["code"] == "instance_lock_lost"


def test_second_backend_rejected_before_recovery(app, monkeypatch):
    from fastapi.testclient import TestClient
    from backend.app import create_app
    monkeypatch.setattr("backend.app._try_build_retrieval_stack", lambda *args: False)
    with TestClient(app) as client:
        with pytest.raises(RuntimeError, match="已有后端"):
            with TestClient(create_app()):
                pass
        assert client.get("/api/health").status_code == 200


def test_binary_parser_timeout_has_safe_error(tmp_path, monkeypatch):
    import subprocess
    path = tmp_path / "slow.pdf"
    path.write_bytes(b"%PDF-1.4 synthetic fixture")
    def timeout(*args, **kwargs):
        raise subprocess.TimeoutExpired("parser", 30)
    monkeypatch.setattr(parsers.subprocess, "run", timeout)
    with pytest.raises(parsers.DocumentParseError, match="document_parse_timeout"):
        parsers.parse_source(path)


def test_post_write_database_failure_removes_source(app, monkeypatch, tmp_path):
    import asyncio
    from io import BytesIO
    from fastapi import UploadFile
    from fastapi.testclient import TestClient
    from sqlalchemy.exc import SQLAlchemyError
    from backend.services.material_service import create_material_record
    monkeypatch.setattr("backend.app._try_build_retrieval_stack", lambda *args: False)
    with TestClient(app):
        with app.state.session_factory() as db:
            original = db.flush
            calls = 0
            def fail(*args, **kwargs):
                nonlocal calls
                calls += 1
                if calls == 2:
                    raise SQLAlchemyError("synthetic post-write failure")
                return original(*args, **kwargs)
            monkeypatch.setattr(db, "flush", fail)
            with pytest.raises(SQLAlchemyError):
                asyncio.run(create_material_record(db, upload=UploadFile(filename="fixture.md", file=BytesIO(b"# Fixture")), title="rollback fixture", materials_root=tmp_path))
            assert list(tmp_path.rglob("source.md")) == []


def test_commit_failure_removes_uncommitted_source(app, monkeypatch, tmp_path):
    from fastapi.testclient import TestClient
    from backend.api.deps import get_db
    from sqlalchemy.exc import SQLAlchemyError
    monkeypatch.setattr("backend.app._try_build_retrieval_stack", lambda *args: False)
    with TestClient(app, raise_server_exceptions=False) as client:
        app.state.settings = app.state.settings.model_copy(update={"upload_dir": tmp_path})
        app.state.job_runner = type("Runner", (), {"notify": lambda self: None, "stop": lambda self: None})()
        with app.state.session_factory() as db:
            app.dependency_overrides[get_db] = lambda: db
            def fail():
                raise SQLAlchemyError("synthetic commit failure")
            monkeypatch.setattr(db, "commit", fail)
            response = client.post("/api/materials?title=commit-fixture", files={"file": ("fixture.md", b"# Fixture", "text/markdown")})
            assert response.status_code == 500
            assert list(tmp_path.rglob("source.md")) == []
        app.dependency_overrides.clear()


def test_docx_expanded_zip_limit(tmp_path, monkeypatch):
    from docx import Document
    path = tmp_path / "expanded.docx"
    Document().save(path)
    monkeypatch.setattr(parsers, "MAX_DOCX_EXPANDED_BYTES", 32)
    with pytest.raises(parsers.DocumentParseError, match="document_content_too_large"):
        parsers.parse_docx(path)


def test_dev_host_rebinding_rejected(app, monkeypatch):
    from fastapi.testclient import TestClient
    monkeypatch.setattr("backend.app._try_build_retrieval_stack", lambda *args: False)
    with TestClient(app, base_url="http://127.0.0.1") as client:
        app.state.settings = app.state.settings.model_copy(update={"app_env": "dev"})
        assert client.get("/api/health", headers={"Host": "untrusted.example"}).status_code == 403
        assert client.get("/api/health").status_code == 200
